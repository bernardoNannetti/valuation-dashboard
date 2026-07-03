"""
Exportação do DCF para Excel - Etapa 3 (final)
=================================================
Pega o resultado de dcf.calcular_dcf() para uma ação e gera um Excel bem
documentado em /valuations/{ticker}.xlsx: premissas usadas, detalhamento do
WACC, projeção de receita/FCFF ano a ano, e o resumo do valuation
(valor da empresa, valor do equity, preço-alvo, recomendação).

Por que os valores vêm hardcoded (não fórmulas do zero como no
comps_analysis)? Aqui a "fonte da verdade" do cálculo é o motor em Python
(src/dcf.py) — já testado, documentado e versionado no Git. Esta planilha é
o RELATÓRIO de saída do modelo, não um modelo vivo em Excel. Ainda assim,
os totais/percentuais que fazem sentido como fórmula (soma dos valores
presentes, upside %) ficam como fórmula, para dar transparência de auditoria.
"""

import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from . import config
from . import dcf

AZUL_ESCURO = "1F4E79"
AZUL_CLARO = "D9E1F2"
CINZA_CLARO = "F2F2F2"
VERDE = "C6EFCE"
VERMELHO = "FFC7CE"
CINZA_NEUTRO = "E7E6E6"
BRANCO = "FFFFFF"

FONTE_TITULO = Font(size=14, bold=True, color=BRANCO)
FONTE_SECAO = Font(size=12, bold=True, color=BRANCO)
FONTE_LABEL = Font(bold=True)
CENTRO = Alignment(horizontal="center", vertical="center")


def _secao(ws, linha, texto, colunas):
    ws.cell(row=linha, column=1, value=texto)
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=colunas)
    for c in range(1, colunas + 1):
        cel = ws.cell(row=linha, column=c)
        cel.fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    ws.cell(row=linha, column=1).font = FONTE_SECAO
    ws.cell(row=linha, column=1).alignment = CENTRO


def _linha_dado(ws, linha, label, valor, formato=None, destaque=None):
    ws.cell(row=linha, column=1, value=label).font = FONTE_LABEL
    c = ws.cell(row=linha, column=2, value=valor)
    if formato:
        c.number_format = formato
    if destaque:
        for col in (1, 2):
            ws.cell(row=linha, column=col).fill = PatternFill("solid", fgColor=destaque)


def gerar_excel_valuation(ticker: str, conn=None, pasta_saida: str = None) -> str:
    fechar = False
    if conn is None:
        conn = sqlite3.connect(config.CAMINHO_BANCO_SQLITE)
        fechar = True
    if pasta_saida is None:
        pasta_saida = config.DIR_VALUATIONS

    r = dcf.calcular_dcf(ticker, conn)
    w = r["wacc_detalhe"]

    if fechar:
        conn.close()

    wb = Workbook()
    ws = wb.active
    ws.title = "DCF"

    total_colunas = 7

    ws.cell(row=1, column=1, value=f"{r['nome']} ({r['ticker']}) — DCF VALUATION")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_colunas)
    ws["A1"].font = FONTE_TITULO
    ws["A1"].fill = PatternFill("solid", fgColor=AZUL_ESCURO)
    ws["A1"].alignment = CENTRO

    ws.cell(row=2, column=1,
            value=f"{r['setor']} / {r['industria']} | As of 03/07/2026 | Valores em USD, exceto preço por ação")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_colunas)
    ws["A2"].alignment = CENTRO
    ws["A2"].font = Font(italic=True, size=9)

    # ---- Resumo do valuation (topo, é o que mais importa) ----------------
    linha = 4
    _secao(ws, linha, "RESUMO DO VALUATION", total_colunas)
    linha += 1
    _linha_dado(ws, linha, "Preço atual", r["preco_atual"], "$#,##0.00"); linha += 1
    _linha_dado(ws, linha, "Preço-alvo (DCF)", r["preco_alvo"], "$#,##0.00"); linha += 1
    linha_upside = linha
    ws.cell(row=linha, column=1, value="Upside / Downside").font = FONTE_LABEL
    ws.cell(row=linha, column=2, value=f"=(B{linha-1}/B{linha-2})-1").number_format = "0.0%"
    linha += 1

    cor_recomendacao = {"Compra": VERDE, "Venda": VERMELHO, "Neutro": CINZA_NEUTRO}.get(r["recomendacao"], CINZA_NEUTRO)
    _linha_dado(ws, linha, "Recomendação", r["recomendacao"], destaque=cor_recomendacao); linha += 1
    _linha_dado(ws, linha, "Threshold Compra / Venda", f">{config.THRESHOLD_COMPRA:.0%} / <{config.THRESHOLD_VENDA:.0%}"); linha += 2

    # ---- Premissas macro e WACC -------------------------------------------
    _secao(ws, linha, "WACC (CAPM)", total_colunas)
    linha += 1
    _linha_dado(ws, linha, "Beta", round(w["beta"], 3)); linha += 1
    _linha_dado(ws, linha, "Taxa livre de risco (Treasury 10Y)", w["taxa_livre_de_risco"], "0.00%"); linha += 1
    _linha_dado(ws, linha, "Prêmio de risco de mercado (ERP, Damodaran)", w["premio_risco_mercado"], "0.00%"); linha += 1
    _linha_dado(ws, linha, "Custo do capital próprio (CAPM)", w["custo_capital_proprio"], "0.00%"); linha += 1
    _linha_dado(ws, linha, "Dívida total", w["divida_total"], "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Custo da dívida pré-imposto (Desp. Juros / Dívida)", w["custo_divida_pre_imposto"], "0.00%"); linha += 1
    _linha_dado(ws, linha, "Taxa de imposto efetiva", w["taxa_imposto_efetiva"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "Custo da dívida pós-imposto", w["custo_divida_pos_imposto"], "0.00%"); linha += 1
    _linha_dado(ws, linha, "Valor de mercado do equity", w["valor_mercado_equity"], "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Peso equity / Peso dívida", f"{w['peso_equity']:.1%} / {w['peso_divida']:.1%}"); linha += 1
    _linha_dado(ws, linha, "WACC", w["wacc"], "0.00%", destaque=AZUL_CLARO); linha += 2

    # ---- Premissas de projeção --------------------------------------------
    _secao(ws, linha, "PREMISSAS DE PROJEÇÃO", total_colunas)
    linha += 1
    _linha_dado(ws, linha, "Horizonte de projeção explícita", f"{r['horizonte_projecao_anos']} anos"); linha += 1
    _linha_dado(ws, linha, "Receita base (mais recente)", r["receita_base"], "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Crescimento de receita — ano 1 (consenso)", r["crescimento_ano1"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "Crescimento na perpetuidade (Gordon Growth)", r["crescimento_perpetuidade"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "Margem EBIT (média histórica)", r["margem_ebit"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "D&A (% da receita, média histórica)", r["pct_da"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "CapEx ano 1 (guidance pública / histórico)", r["pct_capex_inicial"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "CapEx normalizado (média histórica)", r["pct_capex_historico"], "0.0%"); linha += 1
    _linha_dado(ws, linha, "Variação de capital de giro (% da receita)", r["pct_variacao_wc"], "0.0%"); linha += 2

    # ---- Projeção ano a ano -------------------------------------------------
    _secao(ws, linha, "PROJEÇÃO DE FLUXO DE CAIXA (FCFF)", total_colunas)
    linha += 1
    cabecalhos = ["Ano", "Receita", "CapEx (% receita)", "FCFF", "Valor Presente do FCFF"]
    for i, c in enumerate(cabecalhos, start=1):
        cel = ws.cell(row=linha, column=i, value=c)
        cel.font = FONTE_LABEL
        cel.fill = PatternFill("solid", fgColor=AZUL_CLARO)
        cel.alignment = CENTRO
    linha_cab_projecao = linha
    linha += 1
    primeira_linha_projecao = linha
    for i in range(r["horizonte_projecao_anos"]):
        ws.cell(row=linha, column=1, value=i + 1)
        ws.cell(row=linha, column=2, value=round(r["receitas_projetadas"][i])).number_format = "$#,##0"
        ws.cell(row=linha, column=3, value=r["pcts_capex_projetados"][i]).number_format = "0.0%"
        ws.cell(row=linha, column=4, value=round(r["fcffs_projetados"][i])).number_format = "$#,##0"
        ws.cell(row=linha, column=5, value=round(r["valores_presentes_fcff"][i])).number_format = "$#,##0"
        for col in range(1, 6):
            ws.cell(row=linha, column=col).alignment = CENTRO
        linha += 1
    ultima_linha_projecao = linha - 1

    linha += 1
    _linha_dado(ws, linha, "Soma dos VPs dos FCFF projetados",
                f"=SUM(E{primeira_linha_projecao}:E{ultima_linha_projecao})", "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Valor terminal (Gordon Growth)", r["valor_terminal"], "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Valor presente do valor terminal", r["valor_presente_terminal"], "$#,##0"); linha += 1
    linha_valor_empresa = linha
    _linha_dado(ws, linha, "Valor da empresa (Enterprise Value)",
                f"=B{linha-3}+B{linha-1}", "$#,##0", destaque=AZUL_CLARO); linha += 1
    _linha_dado(ws, linha, "(-) Dívida líquida", r["divida_liquida"], "$#,##0"); linha += 1
    _linha_dado(ws, linha, "Valor do equity", f"=B{linha_valor_empresa}-B{linha-1}", "$#,##0", destaque=AZUL_CLARO); linha += 1
    _linha_dado(ws, linha, "Ações em circulação", r["shares_outstanding"], "#,##0"); linha += 1
    _linha_dado(ws, linha, "Preço-alvo por ação (=Valor do equity / Ações)",
                f"=B{linha-2}/B{linha-1}", "$#,##0.00", destaque=AZUL_CLARO); linha += 2

    # ---- Notas -----------------------------------------------------------
    _secao(ws, linha, "NOTAS", total_colunas)
    linha += 1
    notas = [
        "Fonte dos dados: pipeline próprio (yfinance -> SQLite via src/data_fetch.py). "
        "Cálculo do DCF: src/dcf.py (versionado no Git).",
        "CapEx: parte da guidance pública de investimento (quando disponível) e normaliza rápido "
        f"(config.ANOS_FADE_CAPEX={config.ANOS_FADE_CAPEX} anos) para a média histórica — reflete o pico "
        "atual de investimento em IA sem assumir uma trajetória de longo prazo que não temos base para prever.",
        "Limitação conhecida: DCF simples de horizonte curto tende a subestimar mega caps de crescimento "
        "(não captura opcionalidade/crescimento além do período explícito). Ver comps_analysis.xlsx para "
        "cross-check via múltiplos de mercado.",
    ]
    for nota in notas:
        ws.cell(row=linha, column=1, value=f"• {nota}")
        ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=total_colunas)
        ws.cell(row=linha, column=1).font = Font(size=9)
        ws.cell(row=linha, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        linha += 1

    larguras = [42, 20, 18, 18, 20, 12, 12]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura

    import os
    os.makedirs(pasta_saida, exist_ok=True)
    caminho = os.path.join(pasta_saida, f"{ticker}.xlsx")
    wb.save(caminho)
    return caminho


def gerar_todos(tickers=config.TICKERS):
    conn = sqlite3.connect(config.CAMINHO_BANCO_SQLITE)
    caminhos = []
    for ticker in tickers:
        caminho = gerar_excel_valuation(ticker, conn)
        print(f"{ticker}: salvo em {caminho}")
        caminhos.append(caminho)
    conn.close()
    return caminhos


if __name__ == "__main__":
    gerar_todos()
